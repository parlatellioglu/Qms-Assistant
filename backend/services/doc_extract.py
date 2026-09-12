# -*- coding: utf-8 -*-
"""Flatten an uploaded document to plain text for the Compliance Checker.

Mirrors the corpus build's extractors (scripts/corpus/build_cmmi_corpus.py) so an
uploaded file reads exactly like a document already in the corpus — paragraphs as
lines, table/sheet rows as ' | '-joined cells, in document order. Kept self-contained
in the backend (no dependency on the scripts/ tree) so it also works inside the
Docker image, where only backend/ is present.
"""
import os


class UnsupportedFileType(ValueError):
    """Raised for a file extension we have no extractor for."""


def _row_line(cells):
    """Join a row's non-empty cells as ' | ' — the shared table/sheet row format."""
    line = " | ".join(c for c in (str(x).strip() for x in cells if x is not None) if c)
    return line if line.strip(" |") else None


def extract_docx(path: str) -> str:
    """Paragraphs as lines, table rows as ' | '-joined cells (document order)."""
    from docx import Document
    from docx.oxml.ns import qn

    d = Document(path)
    parts = []
    ti = 0
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            t = "".join(n.text or "" for n in child.iter(qn("w:t"))).strip()
            if t:
                parts.append(t)
        elif child.tag == qn("w:tbl"):
            tbl = d.tables[ti]
            ti += 1
            for r in tbl.rows:
                line = _row_line([c.text for c in r.cells])
                if line:
                    parts.append(line)
    return "\n".join(parts)


def extract_xlsx(path: str) -> str:
    """Each sheet's non-empty rows as ' | '-joined cells, prefixed by the sheet name."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = [line for r in ws.iter_rows(values_only=True) if (line := _row_line(r))]
        if rows:
            parts.append(f"[{ws.title}]")
            parts.extend(rows)
    return "\n".join(parts)


def extract_pptx(path: str) -> str:
    """Each slide's text frames and tables in order, slides separated by a marker."""
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    slide_parts.append(t)
            if shape.has_table:
                for r in shape.table.rows:
                    line = _row_line([c.text for c in r.cells])
                    if line:
                        slide_parts.append(line)
        if slide_parts:
            parts.append(f"[Slide {i}]")
            parts.extend(slide_parts)
    return "\n".join(parts)


def extract_plain(path: str) -> str:
    """Read a .txt/.md file as UTF-8 text (invalid bytes replaced, not fatal)."""
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        return fh.read()


EXTRACTORS = {
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".pptx": extract_pptx,
    ".txt": extract_plain,
    ".md": extract_plain,
}

SUPPORTED_EXTENSIONS = tuple(EXTRACTORS)


def extract_file(path: str, filename: str | None = None) -> str:
    """Flatten a file to plain text, dispatching on the extension.

    ``filename`` overrides the extension source when ``path`` is a temp file with a
    generic name (as uploads are). Raises ``UnsupportedFileType`` for unknown types.
    """
    ext = os.path.splitext(filename or path)[1].lower()
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        raise UnsupportedFileType(
            f"Desteklenmeyen dosya türü '{ext or '?'}'. "
            f"Desteklenenler: {', '.join(SUPPORTED_EXTENSIONS)}."
        )
    return extractor(path)
