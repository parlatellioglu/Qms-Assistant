#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extract the filled CMMI documents into a curated Python corpus module:
backend/data/cmmi_docs.py

Sources (searched in order):
  data/CMMI_completed/   .docx lifecycle documents
  data/CMMI_xlsx_pptx/   .xlsx test-case sheets and the .pptx SDLC process deck

Each format is flattened to the same plain-text shape (paragraphs as lines,
table/sheet rows as ' | '-joined cells) so retrieval treats them uniformly. This
mirrors backend/data/small_docs.py (a static list of document dicts), so the
dockerized backend can index everything without reading Office files at runtime.

Re-run after changing any source file (xlsx cases are authored by
scripts/populate/fill_cmmi_cases.py):
    python3 scripts/corpus/build_cmmi_corpus.py
"""
import os
from docx import Document
from docx.oxml.ns import qn
import openpyxl
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DATA_DIR = os.path.join(ROOT, "data")
# Source subdirectories under data/, searched in order for each filename.
SRC_DIRS = ["CMMI_completed", "CMMI_xlsx_pptx"]
OUT_PATH = os.path.join(ROOT, "backend", "data", "cmmi_docs.py")

# Per-document metadata. document_type / status / department / language are kept
# within the enums declared in data/small_docs/metadata_schema.json so the
# existing metadata-validation tests keep passing when CMMI docs surface.
META = {
    "CMMI-IT-001_TRF.docx": {
        "id": "CMMI-IT-001", "title": "Teknik Gereksinim Formu (TRF) - Mobil Self-Servis",
        "document_type": "form", "author": "Kişi 1", "created_date": "2025-03-15",
        "tags": ["cmmi", "trf", "teknik gereksinim", "mobil self-servis", "servis seviyesi", "tarife"],
    },
    "CMMI-IT-002_Tech Reco.docx": {
        "id": "CMMI-IT-002", "title": "Teknik Öneri Dokümanı (Tech Reco) - Mobil Self-Servis",
        "document_type": "report", "author": "Kişi 2", "created_date": "2025-03-15",
        "tags": ["cmmi", "teknik öneri", "mimari", "maliyet", "risk", "mobil self-servis"],
    },
    "CMMI-IT-003_PMP.docx": {
        "id": "CMMI-IT-003", "title": "Proje Yönetim Planı (PMP) - Mobil Self-Servis",
        "document_type": "procedure", "author": "Kişi 1", "created_date": "2025-03-15",
        "tags": ["cmmi", "pmp", "proje yönetim planı", "iş kırılım yapısı", "wbs", "iletişim planı"],
    },
    "CMMI-IT-004_BUR.docx": {
        "id": "CMMI-IT-004", "title": "İş Birimi Gereksinimleri (BUR) - Mobil Self-Servis",
        "document_type": "form", "author": "Kişi 6", "created_date": "2025-03-15",
        "tags": ["cmmi", "bur", "iş gereksinimi", "gereksinim", "mobil self-servis"],
    },
    "CMMI-IT-005_SRS.docx": {
        "id": "CMMI-IT-005", "title": "Sistem Gereksinim Spesifikasyonu (SRS) - Mobil Self-Servis",
        "document_type": "procedure", "author": "Kişi 6", "created_date": "2025-03-15",
        "tags": ["cmmi", "srs", "sistem gereksinimi", "fonksiyonel gereksinim", "güvenlik", "mobil self-servis"],
    },
    "CMMI-IT-006_DSAD.docx": {
        "id": "CMMI-IT-006", "title": "Detaylı Sistem Mimarisi ve Tasarımı (DSAD) - Mobil Self-Servis",
        "document_type": "procedure", "author": "Kişi 2", "created_date": "2025-03-15",
        "tags": ["cmmi", "dsad", "mimari", "tasarım", "veritabanı", "ağ", "mobil self-servis"],
    },
    "CMMI-IT-007_SAT Plan.docx": {
        "id": "CMMI-IT-007", "title": "Sistem Kabul Testi Planı (SAT) - Mobil Self-Servis",
        "document_type": "procedure", "author": "Kişi 5", "created_date": "2025-07-01",
        "tags": ["cmmi", "sat", "sistem kabul testi", "test planı", "kabul kriteri", "mobil self-servis"],
    },
    "CMMI-IT-008_UAT Plan.docx": {
        "id": "CMMI-IT-008", "title": "Kullanıcı Kabul Testi Planı (UAT) - Mobil Self-Servis",
        "document_type": "procedure", "author": "Kişi 5", "created_date": "2025-08-01",
        "tags": ["cmmi", "uat", "kullanıcı kabul testi", "test planı", "kabul kriteri", "mobil self-servis"],
    },
    "CMMI-IT-SAT Certificate.docx": {
        "id": "CMMI-IT-SAT", "title": "Sistem Kabul Sertifikası (SAT Certificate) - Mobil Self-Servis",
        "document_type": "report", "author": "Kişi 5", "created_date": "2025-07-15",
        "tags": ["cmmi", "sat", "sertifika", "kabul", "mobil self-servis"],
    },
    "CMMI-IT-UAT Certificate.docx": {
        "id": "CMMI-IT-UAT", "title": "Kullanıcı Kabul Sertifikası (UAT Certificate) - Mobil Self-Servis",
        "document_type": "report", "author": "Kişi 1", "created_date": "2025-08-15",
        "tags": ["cmmi", "uat", "sertifika", "kabul", "üretime alma", "mobil self-servis"],
    },
    "CMMI-IT-009_SAT Cases.xlsx": {
        "id": "CMMI-IT-009", "title": "Sistem Kabul Testi Senaryoları (SAT Cases) - Mobil Self-Servis",
        "document_type": "report", "author": "Kişi 5", "created_date": "2025-07-11",
        "tags": ["cmmi", "sat", "test senaryosu", "test case", "kabul testi", "mobil self-servis"],
    },
    "CMMI-IT-010_UAT Cases.xlsx": {
        "id": "CMMI-IT-010", "title": "Kullanıcı Kabul Testi Senaryoları (UAT Cases) - Mobil Self-Servis",
        "document_type": "report", "author": "Kişi 5", "created_date": "2025-08-12",
        "tags": ["cmmi", "uat", "test senaryosu", "test case", "kabul testi", "mobil self-servis"],
    },
    "SDLC Process_1.pptx": {
        "id": "CMMI-SDLC", "title": "Solutions Delivery Life Cycle (SDLC) - CMMI Aligned",
        "document_type": "procedure", "author": "SEPG", "created_date": "2011-08-01",
        "department": "Kalite", "language": "EN",
        "tags": ["cmmi", "sdlc", "süreç", "process", "yaşam döngüsü", "şablon", "roller"],
    },
    "CMMI-IT_011_Project Close-out Report.docx": {
        "id": "CMMI-IT-011", "title": "Proje Kapanış Raporu (Close-out) - Mobil Self-Servis",
        "document_type": "report", "author": "Kişi 1", "created_date": "2025-09-15",
        "tags": ["cmmi", "kapanış", "close-out", "lessons learned", "tedarikçi değerlendirme", "mobil self-servis"],
    },
}

# Fixed for every CMMI document.
DEPARTMENT = "Mühendislik"   # within metadata_schema enum
STATUS = "published"         # within metadata_schema enum
LANGUAGE = "TR"
VERSION = "1.0"

# Stable ordering matching the CMMI lifecycle.
ORDER = [
    "CMMI-IT-001_TRF.docx",
    "CMMI-IT-002_Tech Reco.docx",
    "CMMI-IT-003_PMP.docx",
    "CMMI-IT-004_BUR.docx",
    "CMMI-IT-005_SRS.docx",
    "CMMI-IT-006_DSAD.docx",
    "CMMI-IT-007_SAT Plan.docx",
    "CMMI-IT-008_UAT Plan.docx",
    "CMMI-IT-SAT Certificate.docx",
    "CMMI-IT-UAT Certificate.docx",
    "CMMI-IT-009_SAT Cases.xlsx",
    "CMMI-IT-010_UAT Cases.xlsx",
    "CMMI-IT_011_Project Close-out Report.docx",
    "SDLC Process_1.pptx",
]

# Fine-grained template identity per file. document_type (form/procedure/report)
# is too coarse to bind compliance rules to — an SRS and a PMP are both
# "procedure". doc_kind names the actual template so Compliance-Checker rules in
# backend/rules/*.yaml can target a document via applies_to.doc_kind. See
# the Compliance Checker's rule binding (see backend/rules/README.md).
DOC_KIND = {
    "CMMI-IT-001_TRF.docx": "TRF",
    "CMMI-IT-002_Tech Reco.docx": "TECH_RECO",
    "CMMI-IT-003_PMP.docx": "PMP",
    "CMMI-IT-004_BUR.docx": "BUR",
    "CMMI-IT-005_SRS.docx": "SRS",
    "CMMI-IT-006_DSAD.docx": "DSAD",
    "CMMI-IT-007_SAT Plan.docx": "SAT_PLAN",
    "CMMI-IT-008_UAT Plan.docx": "UAT_PLAN",
    "CMMI-IT-SAT Certificate.docx": "SAT_CERT",
    "CMMI-IT-UAT Certificate.docx": "UAT_CERT",
    "CMMI-IT-009_SAT Cases.xlsx": "SAT_CASES",
    "CMMI-IT-010_UAT Cases.xlsx": "UAT_CASES",
    "CMMI-IT_011_Project Close-out Report.docx": "CLOSEOUT",
    "SDLC Process_1.pptx": "SDLC",
}


def _row_line(cells):
    """Join a table/sheet row's non-empty cells as ' | ' — the shared row format."""
    line = " | ".join(c for c in (str(x).strip() for x in cells if x is not None) if c)
    return line if line.strip(" |") else None


def extract_docx(path):
    """Paragraphs as lines, table rows as ' | '-joined cells (document order)."""
    d = Document(path)
    parts = []
    ti = 0
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            t = "".join(n.text or "" for n in child.iter(qn("w:t"))).strip()
            if t:
                parts.append(t)
        elif child.tag == qn("w:tbl"):
            tbl = d.tables[ti]
            ti += 1
            for r in tbl.rows:
                line = _row_line([c.text for c in r.cells])
                if line:
                    parts.append(line)
    return "\n".join(parts)


def extract_xlsx(path):
    """Each sheet's non-empty rows as ' | '-joined cells, prefixed by the sheet name."""
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = [line for r in ws.iter_rows(values_only=True) if (line := _row_line(r))]
        if rows:
            parts.append(f"[{ws.title}]")
            parts.extend(rows)
    return "\n".join(parts)


def extract_pptx(path):
    """Each slide's text frames and tables in order, slides separated by a marker."""
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    slide_parts.append(t)
            if shape.has_table:
                for r in shape.table.rows:
                    line = _row_line([c.text for c in r.cells])
                    if line:
                        slide_parts.append(line)
        if slide_parts:
            parts.append(f"[Slide {i}]")
            parts.extend(slide_parts)
    return "\n".join(parts)


EXTRACTORS = {".docx": extract_docx, ".xlsx": extract_xlsx, ".pptx": extract_pptx}


def extract_text(path):
    """Flatten a source file to plain text, dispatching on its extension."""
    ext = os.path.splitext(path)[1].lower()
    try:
        extractor = EXTRACTORS[ext]
    except KeyError:
        raise ValueError(f"Unsupported file type {ext!r}: {path}")
    return extractor(path)


def find_source(fn):
    """Return (abs_path, rel_path-from-data) for fn across SRC_DIRS, first match."""
    for sub in SRC_DIRS:
        path = os.path.join(DATA_DIR, sub, fn)
        if os.path.exists(path):
            return path, f"data/{sub}/{fn}"
    raise FileNotFoundError(f"{fn} not found under {SRC_DIRS}")


def main():
    docs = []
    for fn in ORDER:
        path, rel = find_source(fn)
        meta = META[fn]
        content = extract_text(path)
        # Prepend the title so the title terms participate in retrieval.
        content = f"{meta['title']}\n{content}"
        docs.append({
            "id": meta["id"],
            "doc_kind": DOC_KIND[fn],
            "title": meta["title"],
            # department / language default to the project-wide values but a
            # source (e.g. the English SDLC deck) may override them in META.
            "department": meta.get("department", DEPARTMENT),
            "document_type": meta["document_type"],
            "status": STATUS,
            "version": VERSION,
            "author": meta["author"],
            "created_date": meta["created_date"],
            "language": meta.get("language", LANGUAGE),
            "tags": meta["tags"],
            "content": content,
            "source_file": rel,
        })

    with open(OUT_PATH, "w", encoding="utf-8") as f:
        f.write('"""Auto-generated by scripts/corpus/build_cmmi_corpus.py — do not edit by hand.\n\n')
        f.write("SYNTHETIC DEMO CORPUS. Curated from the CMMI lifecycle documents in data/, which\n")
        f.write("were written for this project: one fictional project (Mobil Self-Servis Uygulaması,\n")
        f.write("CMMI-IT-2025-014) at a fictional organisation, with invented people, identifiers and\n")
        f.write("integrations. No real company documents or data are included. See data/README.md.\n\n")
        f.write('Mirrors the small_docs.py schema."""\n\n')
        f.write("cmmi_documents = [\n")
        for doc in docs:
            f.write("    {\n")
            for key in ["id", "doc_kind", "title", "department", "document_type", "status",
                        "version", "author", "created_date", "language",
                        "tags", "content", "source_file"]:
                f.write(f"        {key!r}: {doc[key]!r},\n")
            f.write("    },\n")
        f.write("]\n")

    print(f"Wrote {len(docs)} documents to {OUT_PATH}")
    for d in docs:
        print(f"  {d['id']:14s} {d['document_type']:10s} {len(d['content']):5d} chars  {d['title']}")


if __name__ == "__main__":
    main()
